"""M1：Excel 导入 —— 数据总量、可追溯性、幂等性、原始语义保留。"""

from __future__ import annotations

import hashlib
from decimal import Decimal

import pytest

from llm_capacity_pricing.importers.excel_importer import (
    ExcelImportError,
    import_excel,
    parse_deployment_strategy,
)
from llm_capacity_pricing.services import (
    EXPECTED_COUNTS,
    import_excel_into_db,
    open_db,
    validate_import,
)
from llm_capacity_pricing.storage.database import BenchmarkRecordRow


def expected_count(key: str) -> int:
    """任务书 §9 的期望条数。"""
    return EXPECTED_COUNTS[key]

# --------------------------------------------------------------------------
# 数据总量（任务书 §9）
# --------------------------------------------------------------------------


def test_excel_counts_match_expected(excel_db):
    """任务书 §9：GPU 价格 = 2，Token 价格 = 3，部署效能 = 32。"""
    report = validate_import(excel_db)
    checks = {c["name"]: c for c in report["checks"]}
    assert checks["count::hardware_price"]["ok"], checks["count::hardware_price"]["detail"]
    assert checks["count::token_price"]["ok"], checks["count::token_price"]["detail"]
    assert checks["count::benchmark_record"]["ok"], checks["count::benchmark_record"]["detail"]

    from llm_capacity_pricing.services import open_db

    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            counts = repo.counts(session).as_dict()
    finally:
        db.dispose()
    assert counts["hardware_price"] == expected_count("hardware_price") == 2
    assert counts["token_price"] == expected_count("token_price") == 3
    assert counts["benchmark_record"] == expected_count("benchmark_record") == 32


def test_parsed_object_counts(excel_path):
    rep = import_excel(excel_path)
    assert len(rep.hardware_prices) == 2
    assert len(rep.token_prices) == 3
    assert len(rep.benchmarks) == 32
    assert rep.sheet_rows == {"GPU价格": 2, "模型价格": 3, "部署效能": 32}


def test_all_validate_checks_pass(excel_db):
    report = validate_import(excel_db)
    failed = [c for c in report["checks"] if not c["ok"]]
    assert not failed, failed


# --------------------------------------------------------------------------
# 可追溯性
# --------------------------------------------------------------------------


def test_record_ids_trace_back_to_sheet_rows(excel_db):
    """唯一记录 ID 可追溯回原始工作表行号（M1 验收）。"""
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            records = repo.list_benchmarks(session)
            evidences = {e.evidence_id: e for e in repo.list_evidence(session)}
    finally:
        db.dispose()

    excel_records = [r for r in records if r.performance_source_kind == "internal_excel"]
    assert len(excel_records) == 32

    rows = set()
    for rec in excel_records:
        sheet, _, row = rec.record_id.removeprefix("excel_").rpartition("_")
        assert sheet == "部署效能", rec.record_id
        rows.add(int(row))
        ev = evidences[rec.evidence_id]
        assert ev.source_sheet == "部署效能"
        assert ev.source_row == int(row)
        assert ev.source_file.endswith(".xlsx")
    assert rows == set(range(2, 34)), "应覆盖 部署效能 第 2~33 行"


def test_evidence_carries_raw_sha256(excel_db, excel_path):
    expected = hashlib.sha256(excel_path.read_bytes()).hexdigest()
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            for ev in repo.list_evidence(session):
                assert ev.raw_sha256 == expected
    finally:
        db.dispose()


def test_original_cell_values_preserved(excel_db, row2_oracle):
    """I/J/K/L/M/N/O 原始值必须逐字保留，不被复算结果覆盖。"""
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            rec = repo.get_benchmark(session, "excel_部署效能_2")
    finally:
        db.dispose()

    assert rec.deployment_strategy_raw == "TP8DCP8EP1 w/ DP Attn w/ Dspark"
    cols = rec.raw_metrics["columns"]
    assert cols["O_data_source"] == "InferenceX"
    # 期望值来自内部 oracle（data/private/，不入库）；公开克隆会跳过该用例
    assert Decimal(str(cols["I_legacy_input_revenue_10k"])) == row2_oracle[
        "input_revenue_10k_per_gpu_year"
    ]
    assert Decimal(str(cols["M_legacy_revenue_cost_ratio"])) == row2_oracle["revenue_cost_ratio"]
    # legacy_* 字段与原始值一致
    assert (
        rec.legacy_total_revenue_10k_per_gpu_year
        == row2_oracle["total_revenue_10k_per_gpu_year"]
    )


# --------------------------------------------------------------------------
# 幂等性（M1 验收）
# --------------------------------------------------------------------------


def test_reimport_is_idempotent(tmp_path, excel_path):
    db = tmp_path / "idem.db"
    first = import_excel_into_db(excel_path, db)
    second = import_excel_into_db(excel_path, db)
    assert second["counts_before"] == first["counts_after"]
    assert second["counts_after"] == first["counts_after"]
    assert second["idempotent"] is True


def test_reimport_does_not_duplicate_benchmarks(tmp_path, excel_path):
    db = tmp_path / "idem2.db"
    import_excel_into_db(excel_path, db)
    import_excel_into_db(excel_path, db)
    import_excel_into_db(excel_path, db)
    our_db, repo = open_db(db)
    try:
        with our_db.session() as session:
            rows = session.query(BenchmarkRecordRow).all()
    finally:
        our_db.dispose()
    assert len(rows) == 32
    assert len({r.record_id for r in rows}) == 32


# --------------------------------------------------------------------------
# 语义边界（任务书 §2 导入规则）
# --------------------------------------------------------------------------


def test_deployment_gpu_count_is_column_c_not_gpus_per_server(excel_db):
    """C 列部署卡数与 GPU价格!B 每服务器卡数不可混为一谈。"""
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            h200_row9 = repo.get_benchmark(session, "excel_部署效能_9")
            b300_row22 = repo.get_benchmark(session, "excel_部署效能_22")
            gpus = {h.gpu_model: h.gpus_per_server for h in repo.list_hardware_prices(session)}
    finally:
        db.dispose()

    assert gpus == {"B300": 8, "H200": 8}
    assert h200_row9.deployment_gpu_count == 32
    assert b300_row22.deployment_gpu_count == 64
    assert "deployment_gpu_count_differs_from_server_gpu_count" in h200_row9.quality_flags


def test_single_user_p90_not_converted_to_tpot(excel_db):
    """E 列只存 single_user_output_tps_p90，不得推导为 TPOT P99。"""
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            rec = repo.get_benchmark(session, "excel_部署效能_2")
    finally:
        db.dispose()
    assert rec.single_user_output_tps_p90 == Decimal("7.888")
    assert rec.tpot_p99_ms is None
    assert rec.tpot_mean_ms is None


def test_cache_hit_rate_kept_raw_and_semantics_unknown(excel_db, row2_oracle):
    """H 列只作旧表定价代理假设，且必须显式标记。"""
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            rec = repo.get_benchmark(session, "excel_部署效能_2")
    finally:
        db.dispose()
    assert rec.cache_hit_rate_raw == row2_oracle["cache_hit_rate_raw"]
    assert rec.cache_metric_semantics == "unknown"
    assert rec.legacy_proxy_for_billable_hit_rate is True
    assert rec.billable_cached_input_ratio is None
    assert rec.physical_cache_hit_rate is None


@pytest.mark.parametrize(
    ("raw", "expect"),
    [
        (
            "TP8DCP8EP1 w/ DP Attn w/ Dspark",
            {"tp": 8, "ep": 1, "dcp": 8, "dp_attention": True, "dspark": True},
        ),
        ("TP8EP32 w/ DP Attn w/ Dspark", {"tp": 8, "ep": 32}),
        ("TP8EP1 w/ MTP", {"tp": 8, "ep": 1, "mtp": True}),
        ("TP4EP4 w/ MTP", {"tp": 4, "ep": 4}),
    ],
)
def test_strategy_parse_is_marked_unverified(raw, expect):
    parsed = parse_deployment_strategy(raw)
    assert parsed is not None
    assert parsed["verified"] is False
    for k, v in expect.items():
        assert parsed[k] == v


def test_pd_split_strategy_detected_but_unverified():
    parsed = parse_deployment_strategy("P16D16分离")
    assert parsed["prefill_p"] == 16 and parsed["decode_d"] == 16
    assert parsed["deployment_mode_hint"] == "prefill_decode_disaggregated"
    assert parsed["verified"] is False
    # 解析结果不得写入已核实的拓扑字段
    assert "tp" not in parsed or parsed.get("verified") is False


def test_model_names_preserved_verbatim(excel_db):
    db, repo = open_db(excel_db)
    try:
        with db.session() as session:
            names = {t.model_name for t in repo.list_token_prices(session)}
            models = {r.model_name for r in repo.list_benchmarks(session)}
    finally:
        db.dispose()
    assert names == {"KIMI-K3", "GLM-5.2/5.3", "DeepSeek-V4-Pro（峰值）"}
    assert models <= names


# --------------------------------------------------------------------------
# 失败行为（任务书 §9）
# --------------------------------------------------------------------------


def test_missing_file_raises(tmp_path):
    with pytest.raises(ExcelImportError):
        import_excel(tmp_path / "not_here.xlsx")


def test_missing_sheet_raises(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "GPU价格"
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(ExcelImportError, match="缺少工作表"):
        import_excel(p)


def test_negative_price_raises(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "GPU价格"
    ws.append(["GPU型号", "卡数", "摊销或租金价格/卡/年"])
    ws.append(["X", 8, -1])
    ws2 = wb.create_sheet("模型价格")
    ws2.append(["模型", "输入缓存命中/百万Token", "输入缓存未命中/百万Token", "输出/百万Token"])
    ws3 = wb.create_sheet("部署效能")
    ws3.append(["GPU型号", "模型", "部署卡数"])
    p = tmp_path / "neg.xlsx"
    wb.save(p)
    with pytest.raises(ExcelImportError, match="负值"):
        import_excel(p)


def test_invalid_gpu_count_raises(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "GPU价格"
    ws.append(["GPU型号", "卡数", "摊销或租金价格/卡/年"])
    ws2 = wb.create_sheet("模型价格")
    ws2.append(["模型", "输入缓存命中/百万Token", "输入缓存未命中/百万Token", "输出/百万Token"])
    ws3 = wb.create_sheet("部署效能")
    ws3.append(["GPU型号", "模型", "部署卡数"])
    ws3.append(["X", "Y", 0])
    p = tmp_path / "badcount.xlsx"
    wb.save(p)
    with pytest.raises(ExcelImportError, match="部署卡数非法"):
        import_excel(p)


def test_unmatched_model_is_recorded_not_guessed(tmp_path):
    """模型价格缺失时只记录，不猜测补齐。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "GPU价格"
    ws.append(["GPU型号", "卡数", "摊销或租金价格/卡/年"])
    ws.append(["EXAMPLE-GPU-A", 8, 10.0])
    ws2 = wb.create_sheet("模型价格")
    ws2.append(["模型", "输入缓存命中/百万Token", "输入缓存未命中/百万Token", "输出/百万Token"])
    ws2.append(["OTHER-MODEL", 1, 10, 50])
    ws3 = wb.create_sheet("部署效能")
    ws3.append(["GPU型号", "模型", "部署卡数", "并发数", "单用户输出Token/s（P90）",
                "单位卡输入吞吐Token/s/卡", "单位卡输出吞吐Token/s/卡", "总缓存命中率",
                "输入营收 万/年/卡", "输出营收 万/年/卡", "营收 万/年/卡", "成本 万/年/卡",
                "营收/成本", "部署策略", "数据来源"])
    ws3.append(["B300", "UNKNOWN-MODEL", 8, 1, 1.0, 100.0, 10.0, 0.5,
                None, None, None, None, None, "TP8", "InferenceX"])
    p = tmp_path / "unmatched.xlsx"
    wb.save(p)

    rep = import_excel(p)
    assert rep.unmatched_models == ["UNKNOWN-MODEL"]
    rec = rep.benchmarks[0]
    assert "no_matching_token_price" in rec.quality_flags
    assert any("token_price" in m for m in rec.missing_fields)
    # 价格字段没有被猜测填入
    assert rec.legacy_total_revenue_10k_per_gpu_year is None
